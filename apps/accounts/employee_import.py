"""Bulk employee import from the legacy Excel export."""

from __future__ import annotations

import datetime as dt
import io
import re
from dataclasses import dataclass

from django.db import transaction
from openpyxl import load_workbook

from apps.accounts.models import Employee
from apps.organization.models import (
    BuClassification,
    BuLead,
    Center,
    CostCode,
    EmployeeGrade,
    EmployeeGroup,
    Estate,
    Lob,
    Location,
    Process,
    Region,
    Subprocess,
)


MODEL_BY_COLUMN = {
    "BU_Lead_ID": (BuLead, "bu_lead_id"),
    "BUClassification_ID": (BuClassification, "bu_classification_id"),
    "Center_ID": (Center, "center_id"),
    "Process_ID": (Process, "process_id"),
    "Costcode_ID": (CostCode, "cost_code_id"),
    "CurrentLocation_ID": (Location, "location_id"),
    "EmployeeGroup_ID": (EmployeeGroup, "employee_group_id"),
    "Grade": (EmployeeGrade, "employee_grade_id"),
    "LOB_ID": (Lob, "lob_id"),
    "Location_ID": (Location, "location_id"),
    "Region_ID": (Region, "region_id"),
    "Subprocess_ID": (Subprocess, "subprocess_id"),
}

RELATION_BY_COLUMN = {
    "BU_Lead_ID": "bu_lead",
    "BUClassification_ID": "bu_classification",
    "Center_ID": "center",
    "Process_ID": "process",
    "Costcode_ID": "cost_code",
    "CurrentLocation_ID": "current_location",
    "EmployeeGroup_ID": "employee_group",
    "Grade": "employee_grade",
    "LOB_ID": "lob",
    "Location_ID": "location",
    "Region_ID": "region",
    "Subprocess_ID": "subprocess",
}

FIELD_BY_COLUMN = {
    "EmpID": "employee_number",
    "Title": "full_name",
    "Designation": "designation",
    "Domain_ID": "domain_name",
    "EmailID": "email",
    "Gender": "gender",
    "Contact_NO": "contact_number",
    "Emp_Status": "employment_status",
    "LWD": "last_working_date",
    "Date_of_Joining": "date_of_joining",
    "ID": "legacy_row_id",
}

TEMPLATE_HEADERS = [*FIELD_BY_COLUMN.keys(), *MODEL_BY_COLUMN.keys(), "Estate", "Estate_Name", "EstateID", "Manager_EmpID", "Supervisor_EmpID"]


@dataclass
class ImportResult:
    created: int
    updated: int
    estates_created: int
    errors: list[dict[str, object]]


def _clean(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _number(value: object) -> int:
    return int(float(_clean(value).replace(",", "")))


def _date(value: object) -> dt.date | None:
    if value in (None, ""):
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = _clean(value)
    for pattern in ("%m/%d/%Y %I:%M %p", "%d-%m-%Y", "%m/%d/%Y", "%Y-%m-%d"):
        try:
            return dt.datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    raise ValueError(f"Invalid date: {text}")


def _normalize(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().casefold()


def _resolve_reference(column: str, value: object):
    text = _clean(value)
    if not text:
        return None
    model, pk_name = MODEL_BY_COLUMN[column]
    try:
        return model.objects.get(legacy_id=_number(text))
    except (ValueError, TypeError, model.DoesNotExist):
        lookup_field = {
            BuLead: "lead_name",
            BuClassification: "classification_name",
            Center: "center_name",
            CostCode: "cost_code",
            EmployeeGrade: "grade_name",
            EmployeeGroup: "group_name",
            Lob: "lob_name",
            Location: "location_name",
            Process: "process_name",
            Region: "region_name",
            Subprocess: "subprocess_name",
        }[model]
        return next(
            (item for item in model.objects.all() if _normalize(str(getattr(item, lookup_field))) == _normalize(text)),
            None,
        )


def _resolve_estate(value: object, *, allow_create: bool = True) -> tuple[Estate | None, bool]:
    text = _clean(value)
    if not text:
        return None, False
    estate = next(
        (item for item in Estate.all_objects.all() if _normalize(item.estate_name) == _normalize(text)),
        None,
    )
    if estate is None:
        try:
            estate = Estate.all_objects.filter(legacy_id=_number(text)).first()
        except (TypeError, ValueError):
            pass
    if estate:
        if not estate.active_flag:
            estate.restore()
        return estate, False
    try:
        _number(text)
    except (TypeError, ValueError):
        pass
    else:
        return None, False
    if not allow_create:
        return None, False
    return Estate.objects.create(estate_name=text), True


def _header_map(headers: list[object]) -> dict[str, int]:
    return {_clean(header): index for index, header in enumerate(headers) if _clean(header)}


@transaction.atomic
def import_employees(file_bytes: bytes) -> ImportResult:
    workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        headers = _header_map(list(next(rows)))
    except StopIteration:
        return ImportResult(0, 0, 0, [{"row": 1, "detail": "The workbook is empty."}])

    missing = [column for column in ("EmpID", "Title") if column not in headers]
    if missing:
        return ImportResult(0, 0, 0, [{"row": 1, "detail": f"Missing required column(s): {', '.join(missing)}."}])

    created = updated = estates_created = 0
    errors: list[dict[str, object]] = []
    pending: list[tuple[int, dict[str, object]]] = []
    for row_number, values in enumerate(rows, start=2):
        raw = {column: values[index] if index < len(values) else None for column, index in headers.items()}
        if not any(value not in (None, "") for value in raw.values()):
            continue
        try:
            employee_number = _clean(raw.get("EmpID"))
            if not employee_number:
                raise ValueError("EmpID is required")
            fields = {field: _clean(raw[column]) for column, field in FIELD_BY_COLUMN.items() if column in raw}
            fields["legacy_row_id"] = _number(raw["ID"]) if _clean(raw.get("ID")) else None
            fields["date_of_joining"] = _date(raw.get("Date_of_Joining"))
            fields["last_working_date"] = _date(raw.get("LWD"))
            for column in MODEL_BY_COLUMN:
                if column in raw:
                    relation_field = RELATION_BY_COLUMN[column]
                    fields[relation_field] = _resolve_reference(column, raw[column])
                    if _clean(raw[column]) and fields[relation_field] is None:
                        raise ValueError(f"Unknown reference in {column}: {_clean(raw[column])}")
            estate_name = raw.get("Estate") or raw.get("Estate_Name")
            estate_id = raw.get("EstateID")
            estate, estate_created = _resolve_estate(estate_name, allow_create=True) if _clean(estate_name) else _resolve_estate(estate_id, allow_create=False)
            if _clean(estate_name) or _clean(estate_id):
                if estate is None:
                    raise ValueError(f"Unknown estate: {_clean(estate_name or estate_id)}")
            fields["estate"] = estate
            manager_number = _clean(raw.get("Manager_EmpID"))
            supervisor_number = _clean(raw.get("Supervisor_EmpID"))
            pending.append((row_number, {"employee_number": employee_number, "fields": fields, "manager_number": manager_number, "supervisor_number": supervisor_number, "estate_created": estate_created}))
        except (TypeError, ValueError) as exc:
            errors.append({"row": row_number, "detail": str(exc)})

    employees_by_number: dict[str, Employee] = {}
    for row_number, item in pending:
        employee, was_created = Employee.objects.update_or_create(
            employee_number=item["employee_number"], defaults=item["fields"]
        )
        employees_by_number[employee.employee_number] = employee
        created += was_created
        updated += not was_created
        estates_created += item["estate_created"]

    for _, item in pending:
        employee = employees_by_number[item["employee_number"]]
        if item["manager_number"]:
            employee.manager_employee = employees_by_number.get(item["manager_number"]) or Employee.objects.filter(employee_number=item["manager_number"]).first()
        if item["supervisor_number"]:
            employee.supervisor_employee = employees_by_number.get(item["supervisor_number"]) or Employee.objects.filter(employee_number=item["supervisor_number"]).first()
        employee.save(update_fields=["manager_employee", "supervisor_employee", "updated_at"])

    return ImportResult(created, updated, estates_created, errors)