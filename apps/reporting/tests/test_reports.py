"""Report requests: build, render, deliver, schedule (Phase 9.3-9.4)."""

import datetime as dt
import io

import pytest
from django.urls import reverse
from django.utils import timezone
from rest_framework.test import APIClient

from apps.accounts.models import UserEstateScope
from apps.crisis.models import CmscMember
from apps.notifications.models import NotificationLog
from apps.reporting import exports, reports, services
from apps.reporting.models import ReportRequest, ReportType, Schedule
from apps.reporting.tasks import run_scheduled_reports

pytestmark = pytest.mark.django_db


@pytest.fixture(autouse=True)
def _media(settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path / "media"


class TestBuilders:
    def test_estate_detail_rows_mirror_the_cost_code_list(self, org, actor, approved_version):
        UserEstateScope.objects.create(user=actor, estate=org["estate"])
        table = reports.build(ReportType.ESTATE_DETAIL, actor, {})
        assert table.columns[:2] == ["Estate", "Cost code"]
        assert len(table.rows) == 1  # scope: Alpha only
        row = dict(zip(table.columns, table.rows[0], strict=True))
        assert (
            row["Cost code"] == "CC-1001"
            and row["BCP status"] == "Approved"
            and row["Version"] == 1
        )
        assert (
            row["Coordinators"] == "Arun Coordinator"
            and row["BU lead email"] == "priya@example.com"
        )

    def test_all_builders_run_for_admin(self, org, actor, approved_version, user_factory):
        admin = user_factory(email="r-admin@example.com", roles=["BCM_ADMIN"])
        CmscMember.objects.create(
            cost_code=org["cost_code"], member_name="x", member_email="x@example.com"
        )
        for report_type in ReportType.values:
            table = reports.build(report_type, admin, {"estate_id": org["estate"].pk})
            assert table.columns and table.title
        summary = reports.build(ReportType.DASHBOARD_SUMMARY, admin, {})
        assert ["Cost codes", 2] in summary.rows

    def test_every_format_renders_the_same_rows(self, org, actor, approved_version):
        table = reports.build(ReportType.ESTATE_DETAIL, actor, {})
        table.rows = [
            [
                "Alpha Estate",
                "CC-1001",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                1,
                "Approved",
                "",
                dt.date(2026, 9, 1),
                "",
                "",
                "",
            ]
        ]
        xlsx, xlsx_mime = exports.render(table, "xlsx")
        csv_bytes, csv_mime = exports.render(table, "csv")
        pdf, pdf_mime = exports.render(table, "pdf")
        assert (
            xlsx_mime.endswith("sheet") and csv_mime == "text/csv" and pdf_mime == "application/pdf"
        )
        from openpyxl import load_workbook

        sheet = load_workbook(io.BytesIO(xlsx)).active
        header_row = next(i for i in range(1, 6) if sheet.cell(i, 1).value == "Estate")
        assert [c.value for c in sheet[header_row]][:2] == ["Estate", "Cost code"]
        assert sheet.cell(header_row + 1, 2).value == "CC-1001"
        assert sheet.cell(header_row + 1, 14).value == dt.datetime(2026, 9, 1)
        lines = csv_bytes.decode("utf-8-sig").splitlines()
        assert (
            lines[0].startswith("Estate,Cost code")
            and "CC-1001" in lines[1]
            and "2026-09-01" in lines[1]
        )
        assert pdf.startswith(b"%PDF") and exports.render(table, "pdf")[0] == pdf


class TestRequests:
    def test_request_runs_stores_and_emails_a_download_link(
        self, org, actor, approved_version, coordinator_client, api_client
    ):
        response = coordinator_client.post(
            reverse("reporting:report-requests"),
            {
                "report_type": "ESTATE_DETAIL",
                "report_format": "xlsx",
                "estate_id": org["estate"].pk,
            },
            format="json",
        )
        assert response.status_code == 201, response.data
        body = response.data
        assert (
            body["status"] == "COMPLETED"
            and body["run_count"] == 1
            and body["active_flag"] is False
        )
        assert body["document"]["file_name"].startswith("estate_detail-") and body["document"][
            "file_name"
        ].endswith(".xlsx")

        mail = NotificationLog.objects.get(event_type="REPORT_READY")
        assert mail.to_email == "arun.coordinator@example.com" and mail.status == "SENT"
        assert (
            "/documents/" in mail.context["download_url"]
            and "token=" in mail.context["download_url"]
        )

        # The link works without a session; the file is a real workbook.
        path = mail.context["download_url"].split("/api/v1", 1)[1]
        download = api_client.get(f"/api/v1{path}")
        assert download.status_code == 200
        assert b"".join(download.streaming_content)[:2] == b"PK"

        # Another user cannot mint a link for it.
        other = APIClient()
        from django.contrib.auth import get_user_model

        other.force_authenticate(
            user=get_user_model().objects.create_user(
                email="o@example.com", password="x", display_name="o"
            )
        )
        assert (
            other.get(
                reverse("documents:document-link", args=[body["document"]["entity_document_id"]])
            ).status_code
            == 404
        )

    def test_estate_outside_scope_is_refused(
        self, org, actor, approved_version, coordinator_client
    ):
        response = coordinator_client.post(
            reverse("reporting:report-requests"),
            {
                "report_type": "ESTATE_DETAIL",
                "report_format": "csv",
                "estate_id": org["other_estate"].pk,
            },
            format="json",
        )
        assert response.status_code == 400 and "estate_id" in response.data["field_errors"]

    def test_inline_estate_detail_download(self, org, actor, approved_version, coordinator_client):
        response = coordinator_client.get(
            reverse("reporting:estate-detail-report"), {"file_format": "csv"}
        )
        assert response.status_code == 200 and response["Content-Type"] == "text/csv"
        assert b"CC-1001" in response.content
        assert (
            coordinator_client.get(
                reverse("reporting:estate-detail-report"), {"file_format": "exe"}
            ).status_code
            == 400
        )

    def test_schedule_advances_and_the_sweep_runs_due_requests(
        self, org, actor, approved_version, coordinator_client
    ):
        response = coordinator_client.post(
            reverse("reporting:report-requests"),
            {"report_type": "EXEMPTION_REGISTER", "report_format": "csv", "schedule": "DAILY"},
            format="json",
        )
        assert response.status_code == 201
        request = ReportRequest.objects.get(pk=response.data["report_request_id"])
        assert request.run_count == 1 and request.active_flag is True
        assert request.next_run_at is not None and request.next_run_at > timezone.now()

        # Not due yet: the sweep does nothing.
        assert run_scheduled_reports() == 0
        request.next_run_at = timezone.now() - dt.timedelta(minutes=1)
        request.save(update_fields=["next_run_at"])
        assert run_scheduled_reports() == 1
        request.refresh_from_db()
        assert request.run_count == 2 and request.next_run_at > timezone.now()
        assert (
            NotificationLog.objects.filter(event_type="REPORT_READY", entity_id=request.pk).count()
            == 2
        )

        stopped = coordinator_client.post(
            reverse("reporting:report-request-stop", args=[request.pk])
        )
        assert stopped.data["active_flag"] is False and stopped.data["next_run_at"] is None
        assert run_scheduled_reports() == 0

    def test_failure_is_recorded_not_raised(
        self, org, actor, approved_version, coordinator_client, monkeypatch
    ):
        monkeypatch.setattr(
            reports, "build", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("boom"))
        )
        response = coordinator_client.post(
            reverse("reporting:report-requests"),
            {"report_type": "CALL_TREE_RUNS", "report_format": "xlsx"},
            format="json",
        )
        assert response.status_code == 201
        assert response.data["status"] == "FAILED" and "boom" in response.data["last_error"]

    def test_deactivated_requester_produces_nothing(self, org, actor, approved_version):
        request = services.create_request(
            user=actor,
            report_type=ReportType.ESTATE_DETAIL,
            report_format="csv",
            parameters={},
            schedule=Schedule.DAILY,
        )
        actor.is_active = False
        actor.save(update_fields=["is_active"])
        services.run_request(request.pk)
        request.refresh_from_db()
        assert request.status == "FAILED" and "no longer active" in request.last_error

    def test_types_endpoint_lists_formats(self, coordinator_client):
        data = coordinator_client.get(reverse("reporting:report-types")).data
        assert {t["code"] for t in data["types"]} == set(ReportType.values)
        assert [s["code"] for s in data["schedules"]] == ["ONCE", "DAILY", "WEEKLY", "MONTHLY"]

    def test_monthly_next_run_clamps_to_month_length(self):
        jan31 = dt.datetime(2026, 1, 31, 9, 0, tzinfo=dt.UTC)
        assert services.next_run_after(Schedule.MONTHLY, jan31) == dt.datetime(
            2026, 2, 28, 9, 0, tzinfo=dt.UTC
        )
        assert services.next_run_after(Schedule.ONCE, jan31) is None
